"""
PhishDetect - Utility Module
Handles model loading, prediction, LIME explanations, PDF generation, sender analysis.
"""

import re
import html
import numpy as np
import joblib
import os
from datetime import datetime

# Optional file readers
try:
    import PyPDF2
    PDF_OK = True
except ImportError:
    PDF_OK = False

try:
    import docx
    DOCX_OK = True
except ImportError:
    DOCX_OK = False

from lime.lime_text import LimeTextExplainer

# Cache loaded components
_model = None
_vectorizer = None
_lime_comps = None

def load_components():
    """Load model, vectorizer, LIME components (cached)."""
    global _model, _vectorizer, _lime_comps
    if _model is None:
        base_dir = os.path.dirname(os.path.abspath(__file__))
        _model = joblib.load(os.path.join(base_dir, 'model.pkl'))
        _vectorizer = joblib.load(os.path.join(base_dir, 'vectorizer.pkl'))
        raw = joblib.load(os.path.join(base_dir, 'lime_components.pkl'))
        _lime_comps = raw if isinstance(raw, dict) else {}
    return _model, _vectorizer, _lime_comps

# ============================
# TEXT CLEANING
# ============================

def clean_text(text):
    """Standardized email cleaning."""
    if not isinstance(text, str):
        return ""
    text = text.lower()
    text = html.unescape(text)
    text = re.sub(r'<[^>]+>', ' ', text)
    text = re.sub(r'http[s]?://\S+', ' ', text)
    text = re.sub(r'www\.\S+', ' ', text)
    text = re.sub(r'\S+@\S+', ' ', text)
    text = re.sub(r'[^a-z\s]', ' ', text)
    text = re.sub(r'\s+', ' ', text).strip()
    return text

# ============================
# FILE EXTRACTION
# ============================

def extract_text_from_file(uploaded_file):
    """Extract text from TXT, PDF, DOCX."""
    ext = uploaded_file.name.split('.')[-1].lower()
    if ext == 'txt':
        return uploaded_file.read().decode('utf-8')
    elif ext == 'pdf' and PDF_OK:
        reader = PyPDF2.PdfReader(uploaded_file)
        return ' '.join([p.extract_text() or '' for p in reader.pages])
    elif ext == 'docx' and DOCX_OK:
        doc = docx.Document(uploaded_file)
        return ' '.join([p.text for p in doc.paragraphs])
    return None

# ============================
# PREDICTION
# ============================

def predict_email(email_text):
    """Return prediction, confidence, risk, all probabilities."""
    model, vec, _ = load_components()
    cleaned = clean_text(email_text)
    X = vec.transform([cleaned])
    pred = model.predict(X)[0]
    proba = model.predict_proba(X)[0]
    conf = float(np.max(proba) * 100)

    risk_map = {
        'legitimate': ('Low', '🟢'),
        'traditional_phishing': ('High', '🔵'),
        'ai_generated_phishing': ('Critical', '🔴')
    }
    risk, emoji = risk_map.get(pred, ('Unknown', '⚪'))

    return {
        'prediction': pred,
        'confidence': conf,
        'risk_level': risk,
        'risk_emoji': emoji,
        'probabilities': dict(zip(model.classes_, proba.tolist()))
    }

# ============================
# LIME EXPLANATION
# ============================

def get_lime_explanation(email_text, num_features=10):
    """Generate LIME word weights. Builds explainer if missing."""
    model, vec, lc = load_components()

    def predict_fn(texts):
        cleaned = [clean_text(t) for t in texts]
        vectors = vec.transform(cleaned)
        return model.predict_proba(vectors)

    if 'explainer' not in lc:
        lc['explainer'] = LimeTextExplainer(class_names=list(model.classes_))

    explainer = lc['explainer']
    cleaned = clean_text(email_text)
    X = vec.transform([cleaned])
    pred = model.predict(X)[0]
    pred_index = list(model.classes_).index(pred)

    exp = explainer.explain_instance(
        email_text,
        predict_fn,
        num_features=num_features,
        num_samples=300,
        labels=(pred_index,)
    )

    return {
        'prediction': pred,
        'word_weights': exp.as_list(label=pred_index)
    }

# ============================
# SENDER ANALYSIS
# ============================

def analyze_sender(email_text):
    """
    Analyze email headers for sender-based red flags.
    Returns a list of warning strings or a clean message.
    """
    flags = []

    free_domains = [
        'gmail.com', 'yahoo.com', 'hotmail.com', 'outlook.com',
        'aol.com', 'mail.com', 'protonmail.com', 'ymail.com'
    ]

    # Extract From field
    from_match = re.search(r'From:\s*.*?<([^>]+)>', email_text, re.IGNORECASE)
    if not from_match:
        from_match = re.search(r'From:\s*(\S+@\S+)', email_text, re.IGNORECASE)

    # Extract Reply-To field
    reply_match = re.search(r'Reply-To:\s*.*?<([^>]+)>', email_text, re.IGNORECASE)
    if not reply_match:
        reply_match = re.search(r'Reply-To:\s*(\S+@\S+)', email_text, re.IGNORECASE)

    # Extract display name vs email mismatch
    display_name_match = re.search(r'From:\s*"?([^"<\n]+)"?\s*<([^>]+)>', email_text, re.IGNORECASE)

    if from_match:
        from_email = from_match.group(1).strip().lower()
        from_domain = from_email.split('@')[-1] if '@' in from_email else ''

        # Check free domain
        if from_domain in free_domains:
            flags.append(f"⚠️ Sender uses a free email domain ({from_domain})")

        # Check Reply-To mismatch
        if reply_match:
            reply_email = reply_match.group(1).strip().lower()
            reply_domain = reply_email.split('@')[-1] if '@' in reply_email else ''
            if from_domain and reply_domain and from_domain != reply_domain:
                flags.append(f"⚠️ Reply-To domain ({reply_domain}) differs from sender domain ({from_domain})")

        # Check display name vs email domain mismatch
        if display_name_match:
            display_name = display_name_match.group(1).strip().lower()
            # If display name contains a known brand but domain doesn't match
            known_brands = ['paypal', 'amazon', 'google', 'microsoft', 'apple', 'bank', 'fedex', 'dhl', 'netflix']
            for brand in known_brands:
                if brand in display_name and brand not in from_domain:
                    flags.append(f"⚠️ Display name mentions '{brand}' but sender domain is '{from_domain}'")
                    break
    else:
        flags.append("ℹ️ No From header detected — paste full email including headers for sender analysis")

    return flags if flags else ["✅ No sender anomalies detected"]


# ============================
# COLOR MAP
# ============================

COLOR_MAP = {
    'legitimate': {'hex': '#2e7d32', 'emoji': '🟢'},
    'traditional_phishing': {'hex': '#1565c0', 'emoji': '🔵'},
    'ai_generated_phishing': {'hex': '#c62828', 'emoji': '🔴'}
}

# ============================
# PDF REPORT
# ============================

from fpdf import FPDF

def hex_to_rgb(hex_color):
    hex_color = hex_color.lstrip('#')
    return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))

def generate_pdf_report(email_text, result, word_weights):
    """Generate a downloadable PDF report."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font('Arial', 'B', 20)
    pdf.set_text_color(136, 14, 79)
    pdf.cell(0, 10, 'PhishDetect - Analysis Report', 0, 1, 'C')
    pdf.ln(5)

    pdf.set_font('Arial', '', 11)
    pdf.set_text_color(0, 0, 0)
    pdf.cell(0, 7, f'Date: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}', 0, 1)
    pdf.ln(5)

    pdf.set_font('Arial', 'B', 14)
    rgb = hex_to_rgb(COLOR_MAP[result['prediction']]['hex'])
    pdf.set_text_color(*rgb)
    pdf.cell(0, 8, f'Prediction: {result["prediction"].upper().replace("_", " ")}', 0, 1)

    pdf.set_text_color(0, 0, 0)
    pdf.set_font('Arial', '', 11)
    pdf.cell(0, 7, f'Confidence: {result["confidence"]:.2f}%', 0, 1)
    pdf.cell(0, 7, f'Risk Level: {result["risk_level"]}', 0, 1)
    pdf.ln(5)

    # Sender analysis in PDF
    sender_flags = analyze_sender(email_text)
    pdf.set_font('Arial', 'B', 12)
    pdf.cell(0, 7, 'Sender Analysis:', 0, 1)
    for flag in sender_flags:
        pdf.set_font('Arial', '', 10)
        clean_flag = flag.encode('latin-1', 'replace').decode('latin-1')
        pdf.cell(0, 6, f'  {clean_flag}', 0, 1)
    pdf.ln(5)

    pdf.set_font('Arial', 'B', 12)
    pdf.cell(0, 7, 'Top Influential Words:', 0, 1)
    for word, weight in word_weights[:10]:
        sign = '+' if weight > 0 else '-'
        pdf.set_font('Arial', '', 10)
        pdf.cell(0, 6, f'  {sign} {word} ({weight:.4f})', 0, 1)

    pdf.ln(10)
    pdf.set_font('Arial', 'I', 8)
    pdf.set_text_color(100, 100, 100)
    pdf.cell(0, 5, 'Automated analysis by PhishDetect.', 0, 1, 'C')

    pdf.output('/tmp/report.pdf')
    return '/tmp/report.pdf'
